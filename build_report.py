"""
Run this MONTHLY, or whenever you need a document to show someone.

It reads the Obligations and Actions sheets from AuditTracker.xlsx and produces
a formatted Word document grouped by theme. This is the artefact you hand to the
Audit Committee, to John, or to a government inspector to demonstrate that NCC
knows its obligations and is working through them.

It does not change the workbook. It only reads from it and writes a new Word file
with today's date in the name.
"""

from __future__ import annotations

import sys
from collections import defaultdict
from datetime import date, datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import load_workbook

from audit_logic import action_coverage_warnings

WORKBOOK = "AuditTracker.xlsx"
NCC_BLUE = RGBColor(0x00, 0x20, 0x5B)
RED = RGBColor(0xC0, 0x39, 0x2B)
AMBER = RGBColor(0xD4, 0x82, 0x0A)


def read_sheet(ws):
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        rows.append(dict(zip(headers, row)))
    return rows


def parse_date(value):
    """Cells may hold datetimes (after dashboard normalisation) or strings."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d %B %Y", "%d %b %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def fmt_date(value):
    parsed = parse_date(value)
    return parsed.strftime("%d %B %Y") if parsed else (str(value) if value else "Not set")


def add_heading(doc, text, size, color=NCC_BLUE, space_before=12):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = color
    return p


def main():
    try:
        wb = load_workbook(WORKBOOK, data_only=True)
    except FileNotFoundError:
        print(f"'{WORKBOOK}' not found. Run setup_tracker.py first.")
        sys.exit(1)

    obligations = read_sheet(wb["Obligations"])
    actions = read_sheet(wb["Actions"])

    # Group actions under their parent obligation
    actions_by_obligation = defaultdict(list)
    for action in actions:
        actions_by_obligation[action.get("Parent Obligation ID")].append(action)

    doc = Document()

    # Title block
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Nottingham City Council")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = NCC_BLUE

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = subtitle.add_run("Audit Reform and Local Government Reorganisation\nObligations Status Report")
    srun.bold = True
    srun.font.size = Pt(14)

    datep = doc.add_paragraph()
    datep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    drun = datep.add_run(f"Prepared {datetime.today().strftime('%d %B %Y')}")
    drun.italic = True
    drun.font.size = Pt(10)
    drun.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # ── Executive summary ────────────────────────────────────────────────────
    today = date.today()
    total = len(obligations)
    complete = sum(1 for o in obligations if o.get("Status") == "Complete")
    high = sum(1 for o in obligations
               if o.get("Risk Rating") == "High" and o.get("Status") != "Complete")

    overdue, due_soon = [], []
    for o in obligations:
        t = parse_date(o.get("Target Date"))
        if t and o.get("Status") != "Complete":
            if t < today:
                overdue.append(o)
            elif (t - today).days <= 30:
                due_soon.append(o)

    evidence_gaps = [
        o for o in obligations
        if o.get("Status") == "Complete" and not o.get("Evidence Location")
    ]
    uncovered = action_coverage_warnings(obligations, actions)

    doc.add_paragraph()
    summary = doc.add_paragraph()
    summary.add_run(
        f"This report sets out {total} obligations arising from the local audit reform "
        f"programme, the transition to the Local Audit Office, and Local Government "
        f"Reorganisation. Of these, {complete} are complete and {high} open obligations "
        f"are rated high risk. {len(overdue)} are past their target date and "
        f"{len(due_soon)} fall due within 30 days. Each obligation is shown with its "
        f"source, owner, current status, the evidence held or required, and the actions "
        f"being taken against it."
    )

    # ── Attention needed ─────────────────────────────────────────────────────
    if overdue or due_soon or evidence_gaps or uncovered:
        add_heading(doc, "Attention needed", 14, space_before=18)
        for o in overdue:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(
                f"OVERDUE — {o.get('Obligation ID', '')} {o.get('Obligation', '')} "
                f"(owner: {o.get('Owner') or 'unassigned'}; "
                f"target {fmt_date(o.get('Target Date'))})"
            )
            run.font.size = Pt(10)
            run.font.color.rgb = RED
        for o in due_soon:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(
                f"DUE SOON — {o.get('Obligation ID', '')} {o.get('Obligation', '')} "
                f"(owner: {o.get('Owner') or 'unassigned'}; "
                f"target {fmt_date(o.get('Target Date'))})"
            )
            run.font.size = Pt(10)
            run.font.color.rgb = AMBER
        for o in evidence_gaps:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(
                f"EVIDENCE GAP — {o.get('Obligation ID', '')} is Complete but no "
                f"Evidence Location is recorded."
            )
            run.font.size = Pt(10)
            run.font.color.rgb = AMBER
        if uncovered:
            p = doc.add_paragraph(style="List Bullet")
            ids = ", ".join(w.split(":")[0] for w in uncovered)
            run = p.add_run(
                f"NO ACTION RECORDED — {ids}: live obligations without a row in "
                f"the Actions tab. Run populate_tracker.py for suggested actions."
            )
            run.font.size = Pt(10)
            run.font.color.rgb = AMBER

    # Group obligations by theme
    by_theme = defaultdict(list)
    for o in obligations:
        by_theme[o.get("Theme") or "Other"].append(o)

    for theme in sorted(by_theme):
        add_heading(doc, theme, 14, space_before=18)
        for o in by_theme[theme]:
            add_heading(doc, f"{o.get('Obligation ID', '')}  {o.get('Obligation', '')}", 11,
                        color=RGBColor(0x22, 0x22, 0x22), space_before=10)

            facts = [
                ("Source", o.get("Source / Trigger")),
                ("Owner", o.get("Owner") or "Not yet assigned"),
                ("Level", o.get("Obligation Level") or "Not classified"),
                ("Risk", o.get("Risk Rating") or "Not rated"),
                ("Target date", fmt_date(o.get("Target Date"))),
                ("Status", o.get("Status") or "Not started"),
                ("Evidence required", o.get("Evidence Required")),
                ("Evidence location", o.get("Evidence Location") or "None recorded yet"),
            ]
            for label, value in facts:
                if value is None:
                    value = ""
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                lab = p.add_run(f"{label}: ")
                lab.bold = True
                lab.font.size = Pt(10)
                val = p.add_run(str(value))
                val.font.size = Pt(10)

            # Linked actions
            linked = actions_by_obligation.get(o.get("Obligation ID"), [])
            if linked:
                ap = doc.add_paragraph()
                ar = ap.add_run("Actions:")
                ar.bold = True
                ar.font.size = Pt(10)
                for action in linked:
                    bullet = doc.add_paragraph(style="List Bullet")
                    status = action.get("Status") or "Not started"
                    pct = action.get("Percent Complete")
                    pct_text = f" ({pct}%)" if pct not in (None, "") else ""
                    due = parse_date(action.get("Due Date"))
                    due_text = f", due {due.strftime('%d %b %Y')}" if due else ""
                    run = bullet.add_run(
                        f"{action.get('Action', '')} — {status}{pct_text}{due_text}"
                    )
                    run.font.size = Pt(10)
                    if due and due < today and status not in ("Complete", "Cancelled"):
                        flag = bullet.add_run("  [OVERDUE]")
                        flag.font.size = Pt(10)
                        flag.bold = True
                        flag.font.color.rgb = RED
                    if status == "Complete" and not action.get("Evidence Link"):
                        flag = bullet.add_run("  [evidence link missing]")
                        flag.font.size = Pt(10)
                        flag.font.color.rgb = AMBER

    out_name = f"NCC_Obligations_Report_{datetime.today().strftime('%Y-%m-%d')}.docx"
    doc.save(out_name)
    print(f"Created '{out_name}' covering {total} obligations.")
    print(f"  Overdue: {len(overdue)}  Due within 30 days: {len(due_soon)}  "
          f"Evidence gaps: {len(evidence_gaps)}  Without actions: {len(uncovered)}")


if __name__ == "__main__":
    main()
