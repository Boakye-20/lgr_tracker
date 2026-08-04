"""
Rebuilds the Dashboard sheet in AuditTracker.xlsx.

Visual theme: Nottingham City Council Corporate Identity Guidelines.
  - Primary accent: NCC lime green (#78BE20)
  - Title / dark text: NCC near-black navy (#231F20)
  - Alert colours: red / amber as needed
  - Clean white background, no gridlines

Run just before showing the tracker to anyone. Close the workbook first.
"""

from __future__ import annotations

import math
import sys
from collections import Counter
from datetime import date, datetime

from openpyxl import load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabel, DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

WORKBOOK = "AuditTracker.xlsx"

# ── NCC brand palette ────────────────────────────────────────────────────────
NCC_GREEN  = "78BE20"   # corporate lime green
NCC_DARK   = "231F20"   # near-black (body text / title)
NCC_NAVY   = "00205B"   # deep navy (secondary accent)
NCC_LGREY  = "F4F4F4"   # tile / row background
NCC_MGREY  = "DDDDDD"   # border / rule
RED        = "C0392B"
AMBER      = "D4820A"
GREEN_OK   = "27AE60"
WHITE      = "FFFFFF"

# ── Fonts ────────────────────────────────────────────────────────────────────
F_TITLE    = Font(name="Calibri", color=NCC_DARK,  bold=True,  size=18)
F_SUBTITLE = Font(name="Calibri", color="777777",  italic=True, size=9)
F_SECTION  = Font(name="Calibri", color=WHITE,     bold=True,  size=10)
F_TILE_NUM = Font(name="Calibri", color=WHITE,     bold=True,  size=26)
F_TILE_LBL = Font(name="Calibri", color=WHITE,     bold=True,  size=8)
F_COL_HDR  = Font(name="Calibri", color=WHITE,     bold=True,  size=9)
F_BODY     = Font(name="Calibri", color=NCC_DARK,  size=9)
F_BODY_B   = Font(name="Calibri", color=NCC_DARK,  bold=True,  size=9)
F_ALERT_R  = Font(name="Calibri", color=RED,       bold=True,  size=9)
F_ALERT_A  = Font(name="Calibri", color=AMBER,     bold=True,  size=9)
F_OK       = Font(name="Calibri", color=GREEN_OK,  italic=True, size=9)

# ── Borders ──────────────────────────────────────────────────────────────────
_THIN  = Side(style="thin",   color=NCC_MGREY)
_MED   = Side(style="medium", color=NCC_GREEN)
BORD_THIN  = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
BORD_BOT_G = Border(bottom=Side(style="medium", color=NCC_GREEN))

# ── Fills ────────────────────────────────────────────────────────────────────
def fill(hex_colour):
    return PatternFill("solid", fgColor=hex_colour)


# ── Layout grid ──────────────────────────────────────────────────────────────
# Everything sits on one 12-column grid (B..M) of equal-width columns, so the
# four KPI tiles (3 columns each) and the three panels / charts (4 columns
# each) always line up on a column boundary and can never overlap.
FIRST_COL  = 2                              # column B
GRID_COLS  = 12
LAST_COL   = FIRST_COL + GRID_COLS - 1      # column M
COL_W      = 14                             # width in characters of every grid column

TILE_SPAN  = 3
PANEL_SPAN = 4
LABEL_SPAN = PANEL_SPAN - 2                 # label 2 cols, value 1 col, gutter 1 col

TILE_ANCHORS  = [FIRST_COL + i * TILE_SPAN  for i in range(4)]
PANEL_ANCHORS = [FIRST_COL + i * PANEL_SPAN for i in range(3)]

# Excel column width -> pixels (Calibri 11), used to size charts to their panel.
COL_PX     = COL_W * 7 + 5
PANEL_PX   = PANEL_SPAN * COL_PX
GUTTER_PX  = 14
CHART_W_CM = (PANEL_PX - GUTTER_PX) * 2.54 / 96
CHART_H_CM = 8.5
BAND_ROW_PT = 15                            # row height reserved beneath a chart


# ── Helpers ──────────────────────────────────────────────────────────────────
def parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d %B %Y", "%d %b %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def read_rows(ws):
    headers = [c.value for c in ws[1]]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        out.append(dict(zip(headers, row)))
    return out


def normalise_dates(ws, col_header):
    headers = [c.value for c in ws[1]]
    if col_header not in headers:
        return
    col = headers.index(col_header) + 1
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, col)
        parsed = parse_date(cell.value)
        if parsed is not None:
            cell.value = parsed
            cell.number_format = "yyyy-mm-dd"


def _col(n):
    return get_column_letter(n)


def _merge(ws, c1, r1, c2, r2):
    ws.merge_cells(f"{_col(c1)}{r1}:{_col(c2)}{r2}")


def _cell(ws, col, row, value=None, *, font=None, fill_hex=None,
          align_h="left", align_v="center", wrap=False, border=None, fmt=None):
    c = ws.cell(row=row, column=col, value=value)
    if font:
        c.font = font
    if fill_hex:
        c.fill = fill(fill_hex)
    c.alignment = Alignment(horizontal=align_h, vertical=align_v,
                            wrap_text=wrap)
    if border:
        c.border = border
    if fmt:
        c.number_format = fmt
    return c


def _block(ws, col, row, span, value=None, **kw):
    """Write a value into a merged run of `span` columns.

    Every cell in the run is styled before merging, otherwise fills and
    borders only render on the top-left cell of the merge.
    """
    for i in range(span):
        _cell(ws, col + i, row, **kw)
    cell = _cell(ws, col, row, value=value, **kw)
    if span > 1:
        _merge(ws, col, row, col + span - 1, row)
    return cell


def draw_tile(ws, col, row, label, value, colour):
    """3-row tile: accent bar / big number / label."""
    _block(ws, col, row,   TILE_SPAN, fill_hex=NCC_GREEN)          # green cap
    _block(ws, col, row+1, TILE_SPAN, value=value, font=F_TILE_NUM,
           fill_hex=colour, align_h="center")
    _block(ws, col, row+2, TILE_SPAN, value=label, font=F_TILE_LBL,
           fill_hex=colour, align_h="center")

    ws.row_dimensions[row].height   = 6
    ws.row_dimensions[row+1].height = 44
    ws.row_dimensions[row+2].height = 18


def section_bar(ws, row, label):
    """Full-width green section header bar."""
    _block(ws, FIRST_COL, row, GRID_COLS, value=label, font=F_SECTION,
           fill_hex=NCC_GREEN)
    ws.row_dimensions[row].height = 20


def panel_title(ws, col, row, label):
    _block(ws, col, row, LABEL_SPAN, value=label, font=F_BODY_B)
    ws.row_dimensions[row].height = 16


def panel_header(ws, col, row, label, value_label):
    """Dark header row for a panel table (label block + value column)."""
    _block(ws, col, row, LABEL_SPAN, value=label, font=F_COL_HDR,
           fill_hex=NCC_NAVY, border=BORD_THIN)
    _cell(ws, col + LABEL_SPAN, row, value=value_label, font=F_COL_HDR,
          fill_hex=NCC_NAVY, align_h="right", border=BORD_THIN)
    ws.row_dimensions[row].height = 18


def panel_row(ws, col, row, label, value, zebra=False):
    bg = NCC_LGREY if zebra else WHITE
    _block(ws, col, row, LABEL_SPAN, value=label, font=F_BODY,
           fill_hex=bg, border=BORD_THIN)
    _cell(ws, col + LABEL_SPAN, row, value=value, font=F_BODY,
          fill_hex=bg, align_h="right", border=BORD_THIN)
    ws.row_dimensions[row].height = 16


def place_chart(ws, chart, col, row):
    """Anchor a chart to a panel, sized so it cannot spill into the next one."""
    chart.width  = CHART_W_CM
    chart.height = CHART_H_CM
    ws.add_chart(chart, f"{_col(col)}{row}")


def reserve_chart_band(ws, start_row):
    """Give the charts their own rows and return the first free row below."""
    rows = math.ceil((CHART_H_CM * 96 / 2.54) / (BAND_ROW_PT * 96 / 72))
    for r in range(start_row, start_row + rows):
        ws.row_dimensions[r].height = BAND_ROW_PT
    return start_row + rows


# ── Main ─────────────────────────────────────────────────────────────────────
def main():
    try:
        wb = load_workbook(WORKBOOK)
    except FileNotFoundError:
        print(f"'{WORKBOOK}' not found. Run setup_tracker.py first.")
        sys.exit(1)
    except PermissionError:
        print(f"Cannot open '{WORKBOOK}'. Close it in Excel and run again.")
        sys.exit(1)

    normalise_dates(wb["Obligations"], "Target Date")
    normalise_dates(wb["Actions"], "Due Date")
    normalise_dates(wb["Actions"], "Start Date")

    obligations = read_rows(wb["Obligations"])
    actions     = read_rows(wb["Actions"])
    intel       = read_rows(wb["Intelligence"])
    today       = date.today()

    # ── Obligation metrics ───────────────────────────────────────────────────
    total_ob    = len(obligations)
    complete_ob = sum(1 for o in obligations if o.get("Status") == "Complete")
    high_risk   = sum(1 for o in obligations
                      if o.get("Risk Rating") == "High"
                      and o.get("Status") != "Complete")
    overdue_ob, due_soon_ob = [], []
    for o in obligations:
        t = parse_date(o.get("Target Date"))
        if t and o.get("Status") != "Complete":
            if t < today:
                overdue_ob.append(o)
            elif (t - today).days <= 30:
                due_soon_ob.append(o)

    status_order = ["Not started", "In progress", "Complete", "Blocked", "Ongoing"]
    status_counts = Counter(o.get("Status") or "Not started" for o in obligations)
    risk_counts   = Counter(o.get("Risk Rating") or "Not rated" for o in obligations)
    theme_counts  = Counter(o.get("Theme") or "Other" for o in obligations)

    # ── Action metrics ───────────────────────────────────────────────────────
    total_ac     = len(actions)
    complete_ac  = [a for a in actions if a.get("Status") == "Complete"]
    # Two different questions, and the second one flattered us badly. Evidence at
    # completion only looks at finished actions, so with almost nothing complete
    # it reported 100% while most of the register carried no evidence at all.
    # Evidence coverage is the honest denominator: every action on the books.
    evid_gaps    = [a for a in complete_ac if not a.get("Evidence Link")]
    evid_done_pct = (round(100 * (len(complete_ac) - len(evid_gaps)) / len(complete_ac))
                     if complete_ac else 100)
    evid_any     = [a for a in actions if (a.get("Evidence Link") or "").strip()]
    evid_pct     = round(100 * len(evid_any) / total_ac) if total_ac else 0
    overdue_ac   = [a for a in actions
                    if parse_date(a.get("Due Date"))
                    and a.get("Status") not in ("Complete", "Cancelled")
                    and parse_date(a.get("Due Date")) < today]
    pct_vals     = [a["Percent Complete"] for a in actions
                    if isinstance(a.get("Percent Complete"), (int, float))]
    avg_prog     = round(sum(pct_vals) / len(pct_vals)) if pct_vals else 0

    # ── Intelligence metrics ─────────────────────────────────────────────────
    intel_total    = len(intel)
    intel_material = sum(1 for i in intel if i.get("Material") == "Yes")
    intel_unrev    = sum(1 for i in intel if not i.get("Reviewed"))
    intel_recent   = sum(1 for i in intel
                         if i.get("Material") == "Yes"
                         and parse_date(i.get("Date Found"))
                         and (today - parse_date(i.get("Date Found"))).days <= 30)

    # ── Rebuild sheet ────────────────────────────────────────────────────────
    if "Dashboard" in wb.sheetnames:
        del wb["Dashboard"]
    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85

    # ── Column widths — one even 12-column grid, plus a left margin ─────────
    ws.column_dimensions["A"].width = 3
    for c in range(FIRST_COL, LAST_COL + 1):
        ws.column_dimensions[_col(c)].width = COL_W

    # ════════════════════════════════════════════════════════════════════════
    # HEADER  (rows 1-5)
    # ════════════════════════════════════════════════════════════════════════
    ws.row_dimensions[1].height = 8

    # Green accent stripe
    for c in range(1, LAST_COL + 1):
        _cell(ws, c, 2, fill_hex=NCC_GREEN)
    ws.row_dimensions[2].height = 4

    _block(ws, FIRST_COL, 3, GRID_COLS,
           value="NCC Audit Reform and LGR Obligations Tracker", font=F_TITLE)
    ws.row_dimensions[3].height = 30

    _block(ws, FIRST_COL, 4, GRID_COLS,
           value=f"Live status snapshot  ·  Refreshed {today.strftime('%d %B %Y')}",
           font=F_SUBTITLE)
    ws.row_dimensions[4].height = 16

    ws.row_dimensions[5].height = 12

    # ════════════════════════════════════════════════════════════════════════
    # KPI TILES  (rows 6-8)
    # ════════════════════════════════════════════════════════════════════════
    tiles = [
        ("OBLIGATIONS",    total_ob,         NCC_NAVY),
        ("OVERDUE",        len(overdue_ob),  RED   if overdue_ob  else GREEN_OK),
        ("DUE IN 30 DAYS", len(due_soon_ob), AMBER if due_soon_ob else GREEN_OK),
        ("COMPLETE",       complete_ob,      GREEN_OK),
    ]
    for col, (lbl, val, colour) in zip(TILE_ANCHORS, tiles):
        draw_tile(ws, col, 6, lbl, val, colour)

    ws.row_dimensions[9].height = 14

    # ════════════════════════════════════════════════════════════════════════
    # OBLIGATIONS OVERVIEW — three panels, tables above their own chart
    # ════════════════════════════════════════════════════════════════════════
    section_bar(ws, 10, "  OBLIGATIONS OVERVIEW")
    ws.row_dimensions[11].height = 8

    p_status, p_risk, p_theme = PANEL_ANCHORS
    HDR_ROW  = 13
    DATA_ROW = HDR_ROW + 1

    status_pairs = [(s, status_counts.get(s, 0)) for s in status_order]
    risk_order   = ["High", "Medium", "Low"]
    risk_pairs   = [(r, risk_counts.get(r, 0)) for r in risk_order]
    theme_pairs  = sorted(theme_counts.items(), key=lambda x: -x[1])

    panels = [
        (p_status, "Obligations by Status", "Status",      status_pairs),
        (p_risk,   "Obligations by Risk",   "Risk Rating", risk_pairs),
        (p_theme,  "Obligations by Theme",  "Theme",       theme_pairs),
    ]
    for col, title, hdr, pairs in panels:
        panel_title(ws, col, 12, title)
        panel_header(ws, col, HDR_ROW, hdr, "Count")
        for i, (name, cnt) in enumerate(pairs):
            panel_row(ws, col, DATA_ROW + i, name, cnt, zebra=(i % 2 == 1))

    # Charts start below the longest table, with a blank row of breathing space
    spacer_row = DATA_ROW + max(len(p[3]) for p in panels)
    ws.row_dimensions[spacer_row].height = 14
    chart_row  = spacer_row + 1

    def _val_col(col):
        return col + LABEL_SPAN

    # Status bar chart — vertical bars, value labels on each bar
    sc = BarChart()
    sc.title    = "Obligations by Status"
    sc.legend   = None
    sc.grouping = "clustered"
    sc.type     = "col"
    last = HDR_ROW + len(status_pairs)
    sc.add_data(Reference(ws, min_col=_val_col(p_status), min_row=HDR_ROW,
                          max_row=last), titles_from_data=True)
    sc.set_categories(Reference(ws, min_col=p_status, min_row=DATA_ROW,
                                max_row=last))
    sc.series[0].graphicalProperties.solidFill = NCC_GREEN
    sc.series[0].graphicalProperties.line.solidFill = NCC_GREEN
    # Ensure both axes render — openpyxl can silently drop them
    sc.x_axis.delete = False
    sc.y_axis.delete = False
    sc.dLbls = DataLabelList()
    sc.dLbls.showVal     = True
    sc.dLbls.showCatName = False
    sc.dLbls.showSerName = False
    place_chart(ws, sc, p_status, chart_row)

    # Risk pie chart — slices labelled with category name + percentage
    pc = PieChart()
    pc.title  = "Obligations by Risk"
    pc.legend = None            # slice labels already name each category
    last = HDR_ROW + len(risk_pairs)
    pc.add_data(Reference(ws, min_col=_val_col(p_risk), min_row=HDR_ROW,
                          max_row=last), titles_from_data=True)
    pc.set_categories(Reference(ws, min_col=p_risk, min_row=DATA_ROW,
                                max_row=last))
    slice_colours = [RED, AMBER, GREEN_OK]
    for idx, hex_c in enumerate(slice_colours):
        pt = DataPoint(idx=idx)
        pt.graphicalProperties.solidFill = hex_c
        pt.graphicalProperties.line.noFill = True
        pc.series[0].dPt.append(pt)
    pc.dLbls = DataLabelList()
    pc.dLbls.showCatName = True
    pc.dLbls.showPercent = True
    pc.dLbls.showVal     = False
    pc.dLbls.showSerName = False
    # Empty categories have no slice, so their "0%" label lands on the title
    for idx, (_, cnt) in enumerate(risk_pairs):
        if not cnt:
            pc.dLbls.dLbl.append(
                DataLabel(idx=idx, showCatName=False, showPercent=False,
                          showVal=False, showSerName=False,
                          showLegendKey=False, showBubbleSize=False))
    place_chart(ws, pc, p_risk, chart_row)

    # Theme bar chart — horizontal so long theme names are fully readable
    tc = BarChart()
    tc.title  = "Obligations by Theme"
    tc.legend = None
    tc.type   = "bar"   # horizontal — categories appear on left Y-axis
    last = HDR_ROW + len(theme_pairs)
    tc.add_data(Reference(ws, min_col=_val_col(p_theme), min_row=HDR_ROW,
                          max_row=last), titles_from_data=True)
    tc.set_categories(Reference(ws, min_col=p_theme, min_row=DATA_ROW,
                                max_row=last))
    tc.series[0].graphicalProperties.solidFill = NCC_NAVY
    tc.series[0].graphicalProperties.line.solidFill = NCC_NAVY
    # Ensure both axes render — openpyxl can silently drop them
    tc.x_axis.delete = False
    tc.y_axis.delete = False
    tc.dLbls = DataLabelList()
    tc.dLbls.showVal     = True
    tc.dLbls.showCatName = False
    tc.dLbls.showSerName = False
    place_chart(ws, tc, p_theme, chart_row)

    # ════════════════════════════════════════════════════════════════════════
    # ACTIONS & INTELLIGENCE — starts below the chart band, never on top of it
    # ════════════════════════════════════════════════════════════════════════
    ai_top = reserve_chart_band(ws, chart_row) + 1
    ws.row_dimensions[ai_top - 1].height = 14
    section_bar(ws, ai_top, "  ACTIONS & INTELLIGENCE FEED")
    ws.row_dimensions[ai_top + 1].height = 8

    action_lines = [
        ("Total actions",      total_ac),
        ("In progress",        sum(1 for a in actions if a.get("Status") == "In progress")),
        ("Complete",           len(complete_ac)),
        ("Overdue",            len(overdue_ac)),
        ("Avg. progress",      f"{avg_prog}%"),
        ("Evidence coverage",  f"{evid_pct}% ({len(evid_any)}/{total_ac})"),
        ("Evidence at completion", f"{evid_done_pct}%"),
    ]
    intel_lines = [
        ("Items captured",         intel_total),
        ("Flagged material",       intel_material),
        ("Material (last 30 days)", intel_recent),
        ("Awaiting review",        intel_unrev),
    ]
    ai_hdr = ai_top + 2
    for col, hdr, lines in ((p_status, "Action Summary",    action_lines),
                            (p_risk,   "Intelligence Feed", intel_lines)):
        panel_header(ws, col, ai_hdr, hdr, "Count")
        for i, (lbl, val) in enumerate(lines):
            panel_row(ws, col, ai_hdr + 1 + i, lbl, val, zebra=(i % 2 == 1))

    # ════════════════════════════════════════════════════════════════════════
    # ATTENTION NEEDED — full-width table so obligation text is readable
    # ════════════════════════════════════════════════════════════════════════
    attn_top = ai_hdr + 1 + max(len(action_lines), len(intel_lines)) + 1
    ws.row_dimensions[attn_top - 1].height = 14
    section_bar(ws, attn_top, "  ATTENTION NEEDED")
    ws.row_dimensions[attn_top + 1].height = 8

    # column spans across the 12-column grid: ID / Obligation / Owner / Target / Flag
    attn_spans = [1, 6, 3, 1, 1]
    hdr_row    = attn_top + 2
    col = FIRST_COL
    for span, lbl in zip(attn_spans,
                         ["ID", "Obligation", "Owner", "Target", "Flag"]):
        _block(ws, col, hdr_row, span, value=lbl, font=F_COL_HDR,
               fill_hex=NCC_NAVY, border=BORD_THIN)
        col += span
    ws.row_dimensions[hdr_row].height = 18

    flagged = [(o, "Overdue") for o in overdue_ob] + \
              [(o, "Due soon") for o in due_soon_ob]

    if not flagged:
        _block(ws, FIRST_COL, hdr_row + 1, GRID_COLS,
               value="Nothing overdue or due within 30 days — all obligations on track.",
               font=F_OK, fill_hex=WHITE, border=BORD_THIN)
        ws.row_dimensions[hdr_row + 1].height = 20
    else:
        for i, (o, why) in enumerate(flagged):
            r = hdr_row + 1 + i
            bg = NCC_LGREY if i % 2 == 1 else WHITE
            target = parse_date(o.get("Target Date"))
            why_font = F_ALERT_R if why == "Overdue" else F_ALERT_A
            text = (o.get("Obligation") or "").strip()
            if len(text) > 160:
                text = text[:159] + "…"

            values = [
                (o.get("Obligation ID", ""), F_BODY_B, "left",   False),
                (text,                       F_BODY,   "left",   True),
                (o.get("Owner") or "Unassigned", F_BODY, "left", True),
                (target.strftime("%d %b %Y") if target else "", F_BODY, "left", False),
                (why,                        why_font, "center", False),
            ]
            col = FIRST_COL
            for span, (val, fnt, align, wrap) in zip(attn_spans, values):
                _block(ws, col, r, span, value=val, font=fnt, fill_hex=bg,
                       border=BORD_THIN, align_h=align, wrap=wrap)
                col += span
            ws.row_dimensions[r].height = 34

    # ── Footer ───────────────────────────────────────────────────────────────
    footer_row = hdr_row + 1 + max(len(flagged), 1) + 1
    ws.row_dimensions[footer_row - 1].height = 10
    ws.row_dimensions[footer_row].height = 4
    for c in range(1, LAST_COL + 1):
        _cell(ws, c, footer_row, fill_hex=NCC_GREEN)

    try:
        wb.save(WORKBOOK)
    except PermissionError:
        print(f"Cannot save '{WORKBOOK}'. Close it in Excel and run again.")
        sys.exit(1)

    print(f"Dashboard refreshed for {today.strftime('%d %B %Y')}.")
    print(f"  Obligations: {total_ob}  Overdue: {len(overdue_ob)}  Due soon: {len(due_soon_ob)}")
    print(f"  Actions: {total_ac}  Overdue: {len(overdue_ac)}  "
          f"Evidence coverage: {evid_pct}% ({len(evid_any)}/{total_ac})  "
          f"at completion: {evid_done_pct}%")
    print(f"  Intelligence: {intel_total} captured, {intel_material} material, {intel_unrev} awaiting review")


if __name__ == "__main__":
    main()
